"""
Step 02 (2021/2022): Count command frequencies across all sessions.

Reads the per-attacker session dictionary, counts how often each distinct
command string appears, and writes the commands sorted by descending frequency.

Input  : DATA_DIR/session_dict.json
Output : DATA_DIR/sorted_cmd_counter.json , DATA_DIR/2-sorted_cmd_counter.xlsx
"""
import os
import json
from collections import Counter

# === Config ===
DATA_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data")

# Read the session_dict.json file
with open(os.path.join(DATA_DIR, 'session_dict.json'), 'r') as file:
    session_dict = json.load(file)

def save_xlsx(sorted_cmd_counter,dir):
    import openpyxl

    # Create a new Excel workbook
    wb = openpyxl.Workbook()
    # Get the default worksheet (sheet1)
    sheet1 = wb.active
    sheet1.title = "Sheet1"
    sheet1.cell(row=1, column=1, value="command")
    sheet1.cell(row=1, column=2, value="counte")
    current_row = 2
    for key, num in sorted_cmd_counter.items():
            sheet1.cell(row=current_row, column=1, value=key)
            sheet1.cell(row=current_row, column=2, value=num)
            current_row += 1

    wb.save(dir)

# Create an empty counter dictionary
cmd_counter = Counter()

# Iterate over each session
for session_id, commands in session_dict.items():
    # Iterate over each command
    for command in commands:
        # Extract the command part (strip the "CMD: " prefix)
        cmd = command[5:]
        # Increment the command's frequency count
        if cmd:
            cmd_counter[cmd] += 1

# Sort from highest to lowest frequency
sorted_cmd_counter = dict(sorted(cmd_counter.items(), key=lambda item: item[1], reverse=True))

# Output the result
print(sorted_cmd_counter)
with open(os.path.join(DATA_DIR, "sorted_cmd_counter.json"),"w") as f:
    json.dump(sorted_cmd_counter,f)

save_xlsx(sorted_cmd_counter, os.path.join(DATA_DIR, "2-sorted_cmd_counter.xlsx"))
