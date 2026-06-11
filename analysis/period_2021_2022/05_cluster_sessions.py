"""
Step 05 (2021/2022): Cluster sessions by their command-pattern sequence.

For each session, converts its command list into a sequence of abstract command
patterns (simple commands -> head command; complex commands -> operator
pattern), then groups all sessions sharing the same pattern sequence.

Input  : DATA_DIR/session_dict.json
Output : DATA_DIR/5-session归类.json , DATA_DIR/session归类.xlsx
"""
import os
import json
import re

# === Config ===
DATA_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data")

with open(os.path.join(DATA_DIR, 'session_dict.json'), 'r') as file:
    session_dict = json.load(file)

def save_xlsx():
    import openpyxl

    # Create a new Excel workbook
    wb = openpyxl.Workbook()
    # Get the default worksheet (sheet1)
    sheet1 = wb.active
    sheet1.title = "Sheet1"
    sheet1.cell(row=1, column=1, value="Session_pattern")
    sheet1.cell(row=1, column=2, value="Session_id")
    sheet1.cell(row=1, column=3, value="Session_comds")
    current_row = 2
    for key, subdict in sessions_pattern_sorted_dict.items():
        for sub_key, sub_value in subdict.items():
            sheet1.cell(row=current_row, column=1, value= "==>>\n".join(key))
            sheet1.cell(row=current_row, column=2, value=sub_key)
            sheet1.cell(row=current_row, column=3, value= "==>>\n".join(sub_value))
            current_row += 1
    wb.save(os.path.join(DATA_DIR, "session归类.xlsx"))

# Extract the pattern of a complex command
def complex_to_pattern(text):
    # Match whitespace, |, &, ;, {, [ and ( using a regular expression
    #pattern = r'\s+|\||&|;|{|}|\[|\]'
    # Use re.findall to extract words and symbols from the text
    #tokens = re.findall(r'[^\s\|&;{}()\[\]\']+', text)
    comd_head = ["&&", "||", "&", ";", "(", ")", "{", "}", "[", "]", "|",
                 'dd', 'do', 'done', 'while', 'if', 'print', 'free', 'wc', 'which', 'exit', 'grep', 'head', 'tftp',
                 'awk',
                 'cp', "chpasswd", 'ping', 'curl', 'ftpget', 'cut', 'uniq',
                 'uname', 'which', 'w', 'crontab', 'top', 'echo', 'sh', 'shell', 'enable', 'system', 'linuxshell',
                 'cat', '/bin/busybox', 'debug', 'cmd', 'config', 'start', 'su', '/ip', 'ifconfig', 'ls', 'rm', 'cd',
                 'runshellcmd', 'mkdir', 'kill', 'wget', 'pwd', 'development', 'start-shell', 'vshell', 'exec',
                 'busybox', 'help', 'passwd', '/tmp/ntpclient', 'lscpu', '/system', 'ps', 'python', 'apt-get',
                 'screenfetch', 'sudo',
                 'https://cdn.discordapp.com/attachments/834709504049414155/836574658668265522/New_Text_Document.sh',
                 'bash', 'clear', 'cdd', 'chmod', 'admin', 'scp',
                 '执行命令', 'sd',
                 'pkill', 'killSTDIN;', 'endScript;', 'nohup', "!", "then", "fi", "id", "nproc", "unset", "history",
                 "export", "more", "whoami", "la", "lockr", "chattr", "\'for", "for",
                 "apt", "timeout", "yum","perl","chsh","df"]

    #tokens = re.split(r'\s+|(>&)|(>>)|(\|\|)|(&&)|(\s&)|([|;>{}()\[\]\'\'])', text)
    tokens = re.split(r'\s+|(>&)|(>>)|(\|\|)|(&&)|(\s&)|([|;>{}()\[\]\'\'])', text)

    result = []

    for sub_word in tokens:
        # remove consecutive duplicate elements
        if result and sub_word == result[-1]:
            continue
        # add non-consecutive-duplicate command types
        else:

            if sub_word and sub_word[:2] == 'if':
                result.append("if")

            if sub_word and sub_word[0] == '-':
                result.append(sub_word)
            if sub_word and sub_word[0] == '$':
                result.append(sub_word)

            if sub_word and sub_word[0] == '<':
                result.append(sub_word)
            if sub_word and sub_word[0] == '+':
                result.append(sub_word)
            # './ninfo',
            if sub_word and sub_word[:2] == './':
                result.append('./')
            elif sub_word in comd_head:
                result.append(sub_word)
            elif sub_word and sub_word[:2] == '>>':
                result.append('>>')
            elif sub_word and sub_word[0] == '>':
                result.append('>')
            if sub_word and sub_word[0] == "\"root:":
                continue

    return ' '.join(result)

def command_to_pattern(command):
    if not re.search(r'[;|&><()]', command):

        # Split the command into words using a regular expression
        words = re.split(r'\s+', command)
        # The first word is taken as the category
        comd_pattern = words[0]
    else:
        comd_pattern = complex_to_pattern(command)
    return  comd_pattern

session_counter = {}

# Create a dictionary to keep track of sessions
sessions = {}

for session_id, commands in session_dict.items():
    # Convert the list of commands into a tuple to use as a key for the sessions dictionary

    session_patterns=[]
    for comd in commands:
        if comd!='END'and comd!='CMD: ':
            session_patterns.append(command_to_pattern(comd[5:]))
    session_patterns= tuple(session_patterns)
    print(f'{commands}\n{session_patterns}')
    if session_patterns in sessions:
        sessions[session_patterns][session_id]= commands
    else:
        sessions[session_patterns] = {session_id: commands}



sessions_pattern_sorted_dict = dict(sorted(sessions.items(), key=lambda item: len(item[1]), reverse=True))

save_xlsx()
with open(os.path.join(DATA_DIR, '5-session归类.json'), "w", encoding='utf-8') as file:
    string_dict = {str(key): value for key, value in sessions_pattern_sorted_dict.items()}
    json.dump(string_dict, file, indent=4)
    file.write('\n')
