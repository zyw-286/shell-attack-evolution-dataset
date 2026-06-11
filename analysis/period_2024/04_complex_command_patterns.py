"""
Step 04 (2024): Tokenize complex commands into operator patterns.

Reads the complex commands, tokenizes each into an abstract pattern of
recognised command heads and shell operators (``custom_split``), and groups the
concrete commands by their pattern. Also produces per-pattern counts.

Input  : DATA_DIR/3-复杂指令.json
Output : DATA_DIR/4-复杂指令类型分析-count.json,
         DATA_DIR/4-复杂指令类型分析-sorted.json,
         DATA_DIR/4-复杂指令分析.xlsx
"""
import os
import re
import json
from operator import itemgetter

# === Config ===
DATA_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data")

with open(os.path.join(DATA_DIR, '3-复杂指令.json'), 'r') as file:
    session_dict = json.load(file)
#token_dict={}>&

def custom_split(text):
    # Match whitespace, |, &, ;, {, [ and ( using a regular expression
    #pattern = r'\s+|\||&|;|{|}|\[|\]'
    # Use re.findall to extract words and symbols from the text
    #tokens = re.findall(r'[^\s\|&;{}()\[\]\']+', text)
    comd_head = ["&&", "||", "&", ";", "(", ")", "{", "}", "[", "]", "|",
                 'dd', 'do', 'done', 'while', 'if', 'print', 'free', 'wc', 'which', 'exit', 'grep', 'head', 'tftp',
                 'awk',
                 'cp', "chpasswd", 'ping', 'curl', 'ftpget', 'cut', 'uniq',
                 'uname', 'which', 'w', 'crontab', 'top', 'echo', 'sh', 'shell', 'enable', 'system', 'linuxshell',
                 'cat', '/bin/busybox', 'debug', 'cmd', 'config', 'start', 'su', '/ip', 'ifconfig', 'ls', 'rm', 'cd',
                 'runshellcmd', 'mkdir', 'kill', 'wget', 'pwd', 'development', 'start-shell', 'vshell', 'exec',
                 'busybox', 'help', 'passwd', '/tmp/ntpclient', 'lscpu', '/system', 'ps', 'python', 'apt-get',
                 'screenfetch', 'sudo',
                 'https://cdn.discordapp.com/attachments/834709504049414155/836574658668265522/New_Text_Document.sh',
                 'bash', 'clear',  'cdd',  'chmod',  'admin', 'scp',
                 '执行命令', 'sd',
                 'pkill','killSTDIN;','endScript;','nohup',"!","then", "fi","id","nproc","unset", "history", "export","more","whoami","la","lockr","chattr","\'for","for",
                 "apt","timeout","yum","perl","chsh","df"]

    tokens = re.split(r'\s+|(>&)|(>>)|(\|\|)|(&&)|(\s&)|([|;>{}()\[\]\'\'])', text)

    result = []

    for sub_word in tokens:
        # remove consecutive duplicate elements
        if result and sub_word==result[-1]:
            continue
        # add non-consecutive-duplicate command types
        else:


            if sub_word and sub_word[:2] =='if':
                result.append("if")

            if sub_word and sub_word[0] =='-':
                result.append(sub_word)
            if sub_word and sub_word[0] =='$':
                result.append(sub_word)


            if sub_word and sub_word[0] =='<':
                result.append(sub_word)
            if sub_word and sub_word[0] =='+':
                result.append(sub_word)
            #'./ninfo',
            if sub_word and sub_word[:2] == './':
                result.append('./')
            elif sub_word in comd_head :
                result.append(sub_word)
            elif  sub_word and sub_word[:2] =='>>':
                result.append('>>')
            elif sub_word and sub_word[0] =='>':
                result.append('>')
            if sub_word and sub_word[0] =="\"root:" :
                continue
            # if sub_word and sub_word[0] =="\"" :
            #     result.append(sub_word)
            # elif sub_word and sub_word[-1] =="\"" :
            #     result.append(sub_word)
            # if sub_word and sub_word[0] =="\'" :
            #     result.append(sub_word)
            # elif sub_word and sub_word[-1] =="\'" :
            #     result.append(sub_word)


    return result


command_pattern_dict={}
command_pattern_count_dict={}

for command  in session_dict:
    count=session_dict[command]
    pattern=tuple(custom_split(command))
    #token_dict[command] = pattern
    print("---"*10)
    print(f"COMMAND:{command}")
    print(pattern)
    if str(pattern) not in command_pattern_dict:
        command_pattern_dict[str(pattern)]={command:count}
    else:
        command_pattern_dict[str(pattern)][command]=count
command_pattern_sorted_dict = dict(sorted(command_pattern_dict.items(), key=lambda item: len(item[1]), reverse=True))

for key in command_pattern_sorted_dict:
    subdict = command_pattern_sorted_dict[key]
    subdict_length = len(subdict)
    subdict_sum = sum(subdict.values())
    command_pattern_count_dict[key] = [subdict_length, subdict_sum]


def save_json():
    #with open('../data/4-复杂指令类型分析-count.json', "w", encoding='utf-8') as file:
    with open(os.path.join(DATA_DIR, '4-复杂指令类型分析-count.json'), "w", encoding='utf-8') as file:

        json.dump(command_pattern_count_dict, file, indent=4)
    with open(os.path.join(DATA_DIR, '4-复杂指令类型分析-sorted.json'), "w", encoding='utf-8') as file:
        json.dump(command_pattern_sorted_dict, file, indent=4)


def save_xlsx():
    import openpyxl

    # Create a new Excel workbook
    wb = openpyxl.Workbook()
    # Get the default worksheet (sheet1)
    sheet1 = wb.active
    sheet1.title = "Sheet1"
    sheet1.cell(row=1, column=1, value="Key")
    sheet1.cell(row=1, column=2, value="Subdict Key")
    sheet1.cell(row=1, column=3, value="Subdict Value")
    current_row = 2
    for key, subdict in command_pattern_sorted_dict.items():
        for sub_key, sub_value in subdict.items():
            sheet1.cell(row=current_row, column=1, value=key)
            sheet1.cell(row=current_row, column=2, value=sub_key)
            sheet1.cell(row=current_row, column=3, value=sub_value)
            current_row += 1

    sheet2 = wb.create_sheet(title="Sheet2")
    for key, values in command_pattern_count_dict.items():
        sheet2.append([key, values[0], values[1]])

    sheet3 = wb.create_sheet(title="Sheet3")
    for key, values in command_pattern_count_dict.items():
        sheet3.append([list(command_pattern_sorted_dict[key].keys())[0], values[1]])

    wb.save(os.path.join(DATA_DIR, "4-复杂指令分析.xlsx"))

save_json()
save_xlsx()
